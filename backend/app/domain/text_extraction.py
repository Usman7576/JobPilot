"""Resume file sniffing and text extraction (PDF via pdfplumber, DOCX via python-docx)."""

from __future__ import annotations

import io
import re
import zipfile
from collections import Counter
from dataclasses import dataclass
from typing import Literal

import pdfplumber
from docx import Document

FileKind = Literal["pdf", "docx"]
MIME_TYPES: dict[FileKind, str] = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}
MAX_TEXT_CHARS = 25_000
MIN_TEXT_CHARS = 200
_MAX_DOCX_UNCOMPRESSED = 50 * 1024 * 1024  # zip-bomb guard


class ExtractionError(Exception):
    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code
        self.message = message


@dataclass(frozen=True)
class ExtractedText:
    text: str
    pages: int


def sniff_kind(data: bytes) -> FileKind | None:
    """Identify the file by content, not by its name or the browser's content type."""
    if data.startswith(b"%PDF-"):
        return "pdf"
    if data.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as zf:
                if "word/document.xml" in zf.namelist():
                    return "docx"
        except zipfile.BadZipFile:
            return None
    return None


def pdf_page_count(data: bytes) -> int:
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            return len(pdf.pages)
    except Exception as exc:
        raise ExtractionError("unreadable_file", "This PDF could not be opened.") from exc


def extract_text(data: bytes, kind: FileKind, max_pages: int) -> ExtractedText:
    pages = _extract_pdf(data, max_pages) if kind == "pdf" else [_extract_docx(data)]
    text = normalize_text(pages)
    if len(text) < MIN_TEXT_CHARS:
        raise ExtractionError(
            "no_text",
            "We couldn't find any text in this file. Scanned PDFs (images of text) aren't "
            "supported yet; export your resume as a text PDF or DOCX.",
        )
    return ExtractedText(text=text[:MAX_TEXT_CHARS], pages=len(pages))


def _extract_pdf(data: bytes, max_pages: int) -> list[str]:
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            if len(pdf.pages) > max_pages:
                raise ExtractionError(
                    "too_many_pages", f"Resumes can be at most {max_pages} pages long."
                )
            return [(page.extract_text(x_tolerance=1.5, y_tolerance=3) or "") for page in pdf.pages]
    except ExtractionError:
        raise
    except Exception as exc:
        raise ExtractionError("unreadable_file", "This PDF could not be read.") from exc


def _extract_docx(data: bytes) -> str:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            if sum(info.file_size for info in zf.infolist()) > _MAX_DOCX_UNCOMPRESSED:
                raise ExtractionError("unreadable_file", "This DOCX file is too large.")
        doc = Document(io.BytesIO(data))
    except ExtractionError:
        raise
    except Exception as exc:
        raise ExtractionError("unreadable_file", "This DOCX file could not be read.") from exc

    lines = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(" | ".join(dict.fromkeys(cells)))  # merged cells repeat
    return "\n".join(lines)


def normalize_text(pages: list[str]) -> str:
    # Lines that repeat on most pages are headers/footers (name, page numbers).
    repeated: set[str] = set()
    if len(pages) > 1:
        counts = Counter(
            line.strip() for page in pages for line in set(page.splitlines()) if line.strip()
        )
        repeated = {line for line, n in counts.items() if n >= max(2, len(pages) - 1)}

    out: list[str] = []
    for page in pages:
        for raw in page.splitlines():
            line = raw.strip()
            if (
                not line
                or line in repeated
                or re.fullmatch(r"(page\s*)?\d+(\s*(/|of)\s*\d+)?", line, re.I)
            ):
                if out and out[-1] != "":
                    out.append("")
                continue
            out.append(re.sub("[ \t\u00a0]+", " ", line))  # spaces, tabs, no-break spaces

    text = "\n".join(out)
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)  # re-join hyphenated line breaks
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()
